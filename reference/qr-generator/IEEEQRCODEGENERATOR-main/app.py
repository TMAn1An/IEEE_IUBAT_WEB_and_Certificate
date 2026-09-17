from flask import Flask, render_template, request, redirect, url_for, flash, send_from_directory, jsonify
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import qrcode
from qrcode.constants import ERROR_CORRECT_M
from pathlib import Path
from datetime import datetime
import secrets
import string
import threading
import re

app = Flask(__name__)
app.secret_key = "change-this-secret-key-before-public-deployment"

CONFERENCE_OPTIONS = [
    "IEEE BECITHCON 2026",
]
EVENT_OPTIONS = [
    "BECITHCON 2026",
]
BASE_DIR = Path(__file__).resolve().parent
EXCEL_DIR = BASE_DIR / "registrations"
LEGACY_EXCEL_FILE = BASE_DIR / "registrations.xlsx"
QR_DIR = BASE_DIR / "static" / "qr"
EXCEL_DIR.mkdir(parents=True, exist_ok=True)
QR_DIR.mkdir(parents=True, exist_ok=True)

excel_lock = threading.Lock()

HEADERS = [
    "SL",
    "Conference",
    "Role",
    "Name",
    "Session",
    "Codeword",
    "Created At",
    "QR File"
]

PRESET_ROLES = [
    "Session Chair",
    "Invited Speaker",
    "Keynote Speaker",
    "Volunteer"
]


def safe_excel_text(value: str) -> str:
    """
    Prevent spreadsheet formula injection if a value starts with =, +, -, or @.
    """
    value = (value or "").strip()
    if value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def filename_part(value: str) -> str:
    value = re.sub(r"[^A-Za-z0-9]+", "_", value or "").strip("_")
    return value[:70] or "Unnamed"


def get_excel_file(conference_type: str, conference: str, role: str) -> Path:
    type_name = conference_type or "No_Conference"
    conference_name = conference or "No_Conference"
    filename = (
        f"{filename_part(type_name)}_{filename_part(conference_name)}_"
        f"Role_{filename_part(role)}.xlsx"
    )
    return EXCEL_DIR / filename


def ensure_excel_file(excel_file: Path):
    if excel_file.exists():
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "Registrations"
    ws.append(HEADERS)

    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = {
        1: 8,
        2: 28,
        3: 22,
        4: 28,
        5: 65,
        6: 24,
        7: 22,
        8: 32,
    }
    for col_idx, width in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A2"
    wb.save(excel_file)


def load_existing_codewords():
    codes = set()
    excel_files = list(EXCEL_DIR.glob("*.xlsx"))
    if LEGACY_EXCEL_FILE.exists():
        excel_files.append(LEGACY_EXCEL_FILE)
    for excel_file in excel_files:
        wb = load_workbook(excel_file, read_only=True)
        ws = wb["Registrations"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[5]:
                codes.add(str(row[5]))
        wb.close()
    return codes


def generate_unique_codeword(length=16):
    alphabet = string.ascii_uppercase + string.digits
    existing = load_existing_codewords()

    while True:
        code = "".join(secrets.choice(alphabet) for _ in range(length))
        if code not in existing:
            return code


def record_exists(excel_file, name, role, session):
    ensure_excel_file(excel_file)
    wb = load_workbook(excel_file, read_only=True)
    ws = wb["Registrations"]

    n = name.strip().casefold()
    r = role.strip().casefold()
    s = session.strip().casefold()

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[2] or not row[3]:
            continue
        if (
            str(row[3]).strip().casefold() == n
            and str(row[2]).strip().casefold() == r
            and str(row[4] or "").strip().casefold() == s
        ):
            codeword = str(row[5]) if row[5] else ""
            wb.close()
            return codeword

    wb.close()
    return None


def qr_payload(conference_type, conference, role, name, session, codeword):
    fields = [f"Role: {role}", f"Name: {name}"]
    if conference:
        fields.insert(0, f"{conference_type}: {conference}")
    if session:
        fields.append(f"Session: {session}")
    fields.append(f"Codeword: {codeword}")
    return "\n".join(fields)


def create_qr(conference_type, conference, role, name, session, codeword):
    payload = qr_payload(conference_type, conference, role, name, session, codeword)

    qr = qrcode.QRCode(
        version=None,
        error_correction=ERROR_CORRECT_M,
        box_size=14,
        border=4,
    )
    qr.add_data(payload)
    qr.make(fit=True)

    img = qr.make_image(fill_color="black", back_color="white")
    filename = f"{codeword}.png"
    filepath = QR_DIR / filename
    img.save(filepath)
    return filename


def append_registration(excel_file, conference, role, name, session, codeword, qr_filename):
    ensure_excel_file(excel_file)

    wb = load_workbook(excel_file)
    ws = wb["Registrations"]
    serial = ws.max_row

    ws.append([
        serial,
        safe_excel_text(conference),
        safe_excel_text(role),
        safe_excel_text(name),
        safe_excel_text(session),
        codeword,
        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        qr_filename,
    ])

    for cell in ws[ws.max_row]:
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    wb.save(excel_file)
    wb.close()


def get_recent_records(limit=10):
    records = []
    excel_files = list(EXCEL_DIR.glob("*.xlsx"))
    if LEGACY_EXCEL_FILE.exists():
        excel_files.append(LEGACY_EXCEL_FILE)
    for excel_file in excel_files:
        wb = load_workbook(excel_file, read_only=True)
        ws = wb["Registrations"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            records.append({
                "sl": row[0],
                "conference": row[1],
                "role": row[2],
                "name": row[3],
                "session": row[4],
                "codeword": row[5],
                "created_at": row[6],
                "qr_file": row[7],
            })
        wb.close()
    records.sort(key=lambda record: str(record["created_at"] or ""), reverse=True)
    return records[:limit]


@app.route("/", methods=["GET"])
def index():
    return render_template(
        "index.html",
        conference_options=CONFERENCE_OPTIONS,
        event_options=EVENT_OPTIONS,
        recent_records=get_recent_records(10),
        result=None,
        preset_roles=PRESET_ROLES
    )


@app.route("/generate", methods=["POST"])
def generate():
    role_select = request.form.get("role_select", "").strip()
    conference_type = request.form.get("conference_type", "").strip()
    conference_select = request.form.get("conference_select", "").strip()
    include_conference = request.form.get("include_conference") == "on"
    include_session = request.form.get("include_session") == "on"
    name = request.form.get("name", "").strip()
    session = request.form.get("session", "").strip()

    role = role_select
    conference = conference_select
    if not include_conference:
        conference = ""
    if not include_session:
        session = ""

    if not role or not name or (include_conference and not conference):
        flash("Role and Name are required. Select a Conference/Event name when that field is included.", "error")
        return redirect(url_for("index"))

    with excel_lock:
        excel_file = get_excel_file(conference_type, conference, role)
        duplicate_code = record_exists(excel_file, name, role, session)
        if duplicate_code:
            flash(
                f"This exact person/role/session entry already exists. Existing codeword: {duplicate_code}",
                "warning"
            )
            return redirect(url_for("index"))

        codeword = generate_unique_codeword(16)
        qr_filename = create_qr(conference_type, conference, role, name, session, codeword)
        append_registration(excel_file, conference, role, name, session, codeword, qr_filename)

    result = {
        "conference": conference,
        "role": role,
        "name": name,
        "session": session,
        "codeword": codeword,
        "qr_filename": qr_filename,
        "excel_filename": excel_file.name,
    }

    return render_template(
        "index.html",
        conference_options=CONFERENCE_OPTIONS,
        event_options=EVENT_OPTIONS,
        recent_records=get_recent_records(10),
        result=result,
        preset_roles=PRESET_ROLES
    )


@app.route("/download-qr/<filename>")
def download_qr(filename):
    return send_from_directory(QR_DIR, filename, as_attachment=True)


@app.route("/download-excel/<filename>")
def download_excel(filename):
    return send_from_directory(EXCEL_DIR, filename, as_attachment=True)


@app.route("/health")
def health():
    return jsonify({"status": "ok", "conference_options": CONFERENCE_OPTIONS, "event_options": EVENT_OPTIONS})


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5000, debug=True)